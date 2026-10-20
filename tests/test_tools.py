import docx
import pytest

from app.ai.ollama_client import OllamaClient, OllamaError
from app.export.docx_export import export_text
from app.storage.index import Chunk
from app.tools import essay as E
from app.tools import quiz as Q


@pytest.fixture
def client(ollama_url):
    return OllamaClient(ollama_url)


CHUNKS = [Chunk("lib:Bio", "Bio notes", "Slide 2", "Mitochondria make ATP.")]


def test_normalize_question():
    mc = Q.normalize_question({"type": "multiple_choice", "question": "q", "choices": ["a", "b", "c"], "answer": "b."})
    assert mc.answer == "b"
    assert Q.normalize_question({"type": "multiple_choice", "question": "q", "choices": ["a", "b"], "answer": "zzz"}) is None
    tf = Q.normalize_question({"type": "true_false", "question": "q", "answer": "False!"})
    assert tf.choices == ["True", "False"] and tf.answer == "False"
    assert Q.normalize_question({"type": "essay", "question": "q", "answer": "a"}) is None


def test_generate_and_grade_quiz(client):
    quiz = Q.generate_quiz(client, "fake-chat", {}, CHUNKS, 3, list(Q.QUESTION_TYPES), "Easy", True)
    assert [q.type for q in quiz.questions] == ["multiple_choice", "true_false", "short_answer"]
    assert quiz.questions[0].answer == "Mitochondria"          # letter "B" mapped to the choice text
    quiz.questions[0].user_answer = "Mitochondria"
    quiz.questions[1].user_answer = "False"
    quiz.questions[2].user_answer = "energy"
    Q.grade_quiz(client, "fake-chat", {}, quiz)
    assert [q.correct for q in quiz.questions] == [True, False, True]
    assert quiz.score == pytest.approx(66.7)
    assert len(quiz.missed()) == 1
    again = Q.Quiz.from_dict(quiz.to_dict())
    assert again.questions[1].correct is False and again.score == quiz.score


def test_grade_unanswered_and_quiz_type_filter(client):
    q = Q.Question("short_answer", "Why?", "Because")
    assert Q.grade_question(client, "fake-chat", {}, q).correct is False
    only_tf = Q.generate_quiz(client, "fake-chat", {}, CHUNKS, 5, ["true_false"], "Hard", False)
    assert [q.type for q in only_tf.questions] == ["true_false"]


def test_flashcards_and_messages(client):
    cards = Q.generate_flashcards(client, "fake-chat", {}, CHUNKS, 5, True)
    assert cards[0] == {"front": "ATP", "back": "Energy currency"}
    msgs = Q.explain_messages("What is ATP?", CHUNKS, True)
    assert "[Bio notes - Slide 2]" in msgs[1]["content"] and "learning mode" in msgs[0]["content"]
    assert "learning mode" not in Q.study_guide_messages(CHUNKS, False)[0]["content"]


def test_split_outline_variants():
    md = "## I. Intro\n- hook\n## II. Body\n- point\n## III. Conclusion\n- wrap"
    assert len(E.split_outline(md)) == 3
    assert len(E.split_outline("I. Intro\n  A. hook\nII. Body\n  A. point")) == 2
    assert E.split_outline("just one blob of text") == ["just one blob of text"]


def test_draft_streams_each_section_with_context(client):
    p = E.EssayProject(prompt="Explain cell energy", outline="## I. Intro\n- hook\n## II. Body\n- ATP\n## III. End\n- wrap")
    queries = []
    text = "".join(E.draft_stream(client, "fake-chat", {}, p, lambda q: queries.append(q) or CHUNKS, True))
    assert text.count("helpful answer") == 3 and "\n\n" in text
    assert len(queries) == 3
    from tests.conftest import FakeOllama
    last = FakeOllama.chat_log[-1]["messages"][1]["content"]
    assert "section 3 of 3" in last and "helpful answer" in last  # earlier sections are passed along


def test_project_roundtrip_and_citation_prompt():
    p = E.EssayProject(title="T", thesis="X", citation_style="APA")
    d = p.to_dict()
    assert E.EssayProject.from_dict({**d, "unknown": 1}).thesis == "X"
    msgs = E.citation_messages(p, [{"source file": "Notes", "author": "Lee"}], False)
    assert "APA" in msgs[1]["content"] and "author: Lee" in msgs[1]["content"]


def test_docx_export(tmp_path):
    export_text("My Essay", "# Heading\nSome **bold** text\n- item\n1. step", tmp_path / "e.docx")
    d = docx.Document(str(tmp_path / "e.docx"))
    assert [p.text for p in d.paragraphs][:2] == ["My Essay", "Heading"]
    assert any(r.bold and r.text == "bold" for p in d.paragraphs for r in p.runs)
    export_text("t", "plain", tmp_path / "e.md")
    assert (tmp_path / "e.md").read_text() == "plain"
