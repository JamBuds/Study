"""Headless smoke tests that drive the real screens against the fake Ollama server."""
import time
from pathlib import Path

import pytest
from PySide6.QtWidgets import QApplication

from app.ai.workers import finish_all_workers
from app.main import MainWindow
from app.storage import classes as C
from app.ui import add_info, class_shell
from app.ui.convert_view import ConvertPage
from app.ui.essay_view import EssayPage
from app.ui.quiz_view import QuizPage


@pytest.fixture(scope="module")
def qapp():
    return QApplication.instance() or QApplication([])


def wait_for(cond, timeout=20):
    end = time.time() + timeout
    while time.time() < end:
        QApplication.processEvents()
        if cond():
            return
        time.sleep(0.01)
    raise AssertionError("timed out waiting for condition")


@pytest.fixture
def shell(qapp, settings, monkeypatch, tmp_path):
    monkeypatch.setattr(add_info.AddInfoDialog, "exec", lambda self: 0)  # skip the modal "add info" prompt
    from PySide6.QtWidgets import QMessageBox
    for fn in ("information", "warning"):  # modal boxes would block a headless run
        monkeypatch.setattr(QMessageBox, fn, staticmethod(lambda *a, **k: QMessageBox.Ok))
    from pptx import Presentation
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[1]).shapes.title.text = "Mitochondria make ATP"
    prs.save(tmp_path / "bio.pptx")
    C.create_class("Biology")
    win = MainWindow(settings)
    win.open_class("Biology")
    win.shell.tmp = tmp_path
    yield win.shell
    finish_all_workers()
    win.close_shell()


def test_picker_lists_classes(qapp, settings):
    win = MainWindow(settings)
    assert win.picker.grid.count() == 1                      # "create your first class" card
    C.create_class("Chem")
    win.picker.refresh()
    assert win.picker.grid.count() == 1 and win.stack.currentWidget() is win.picker
    win.close()


def test_convert_then_essay_flow(shell):
    # convert a pptx and save it to the library
    shell.show_convert()
    page = shell.stack.currentWidget()
    assert isinstance(page, ConvertPage)
    page.add_paths([shell.tmp / "bio.pptx"])
    page.convert()
    wait_for(lambda: page.convert_btn.isEnabled() and page.results)
    assert "Mitochondria make ATP" in next(iter(page.results.values()))
    page.save_library()
    wait_for(lambda: shell.knowledge.index.has_source("lib:bio"))  # file written AND indexed

    # essay: thesis -> outline -> draft -> history
    shell.start_essay()
    essay = shell.stack.currentWidget()
    assert isinstance(essay, EssayPage)
    wait_for(lambda: essay.picker.list.count() == 1)
    assert essay.picker.selected_ids() == ["lib:bio"]
    essay.title.setText("Cell energy")
    essay.prompt.setPlainText("Explain how cells get energy")
    essay.gen_thesis()
    wait_for(lambda: "helpful answer" in essay.thesis.text() and not essay.thesis.busy)
    assert "hidden" not in essay.thesis.text()               # <think> block stripped
    essay.outline.set_text("## I. Intro\n- hook\n## II. Body\n- ATP")
    essay.gen_draft()
    wait_for(lambda: essay.draft.text().count("helpful answer") == 2 and not essay.draft.busy)
    wait_for(lambda: shell.store.list_history("essay"))
    saved = shell.store.load_history("essay", shell.store.list_history("essay")[0]["id"])
    assert saved["title"] == "Cell energy" and "helpful answer" in saved["draft"]

    # reopen from history
    shell.show_history()
    hist = shell.stack.currentWidget()
    assert hist.lists["essay"].count() == 1
    shell.open_history("essay", saved["id"])
    reopened = shell.stack.currentWidget()
    assert isinstance(reopened, EssayPage) and reopened.prompt.toPlainText() == "Explain how cells get energy"
    assert reopened.picker.selected_ids() == ["lib:bio"]


def test_quiz_flow(shell):
    shell.store.add_library_text("notes", "Mitochondria make ATP. Ribosomes make proteins.")
    shell.knowledge.sync_library()
    shell.start_quiz()
    page = shell.stack.currentWidget()
    assert isinstance(page, QuizPage)
    practice = page.practice
    practice.generate()
    wait_for(lambda: len(practice.cards) == 3)
    cards = practice.cards
    cards[0].radios[1].setChecked(True)      # Mitochondria (correct)
    cards[1].radios[1].setChecked(True)      # False (wrong)
    cards[2].text.setPlainText("energy")
    practice.grade()
    wait_for(lambda: practice.quiz.score is not None and practice.retake.isVisibleTo(practice))
    assert practice.quiz.score == pytest.approx(66.7)
    assert practice.review.isVisibleTo(practice)
    practice.review_missed()
    assert len(practice.cards) == 1
    wait_for(lambda: shell.store.list_history("quiz"))

    # explain a question and flashcards
    page.explain.question.setPlainText("What makes ATP?")
    page.explain.explain()
    wait_for(lambda: "helpful answer" in page.explain.out.text() and not page.explain.out.busy)
    page.flash.generate()
    wait_for(lambda: page.flash.cards)
    assert page.flash.face.text() == "ATP"
    page.flash.flip()
    assert page.flash.face.text() == "Energy currency"
    wait_for(lambda: shell.store.list_history("flashcards"))

    # quiz + deck reopen from history
    quiz_id = shell.store.list_history("quiz")[0]["id"]
    shell.open_history("quiz", quiz_id)
    assert len(shell.stack.currentWidget().practice.cards) == 3
    deck_id = shell.store.list_history("flashcards")[0]["id"]
    shell.open_history("flashcards", deck_id)
    assert shell.stack.currentWidget().flash.cards[0]["front"] == "ATP"


def test_dashboard_and_library_pages(shell):
    shell.store.add_library_text("notes", "hello")
    shell.show_library()
    lib = shell.stack.currentWidget()
    wait_for(lambda: lib.list.count() == 1)
    lib.list.setCurrentRow(0)
    assert lib.view.toPlainText() == "hello"
    shell.show_dashboard()
    assert shell.stack.count() == 1
