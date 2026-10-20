import docx
import fitz
import pytest
from pptx import Presentation
from pptx.util import Inches

from app.convert import convert_file
from app.convert.captions import clean_captions


def test_pptx_with_notes_and_table(tmp_path):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Cell Biology"
    s.placeholders[1].text = "Mitochondria make ATP"
    s.notes_slide.notes_text_frame.text = "Remember the Krebs cycle"
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Data"
    tbl = s2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    tbl.cell(0, 0).text, tbl.cell(0, 1).text = "Organelle", "Job"
    tbl.cell(1, 0).text, tbl.cell(1, 1).text = "Ribosome", "Proteins"
    path = tmp_path / "bio.pptx"
    prs.save(path)
    text = convert_file(path)
    assert "[Slide 1]" in text and "[Slide 2]" in text
    assert "Mitochondria make ATP" in text and "Speaker notes:" in text and "Krebs cycle" in text
    assert "Ribosome | Proteins" in text


def test_docx(tmp_path):
    d = docx.Document()
    d.add_paragraph("First paragraph.")
    d.add_paragraph("Second paragraph.")
    d.save(tmp_path / "n.docx")
    assert convert_file(tmp_path / "n.docx") == "First paragraph.\n\nSecond paragraph."


def test_srt_and_vtt_cleanup(tmp_path):
    srt = "1\n00:00:01,000 --> 00:00:03,000\nHello <i>class</i>\n\n2\n00:00:03,000 --> 00:00:05,000\nHello class\n\n3\n00:00:05,000 --> 00:00:07,000\ntoday we study cells\n"
    vtt = "WEBVTT\n\nNOTE this is a comment\n\n00:01.000 --> 00:03.000 align:start\nHello class\n\n00:03.000 --> 00:05.000\ntoday we study cells\n"
    assert clean_captions(srt) == "Hello class today we study cells"
    assert clean_captions(vtt) == "Hello class today we study cells"
    p = tmp_path / "a.srt"
    p.write_text(srt)
    assert convert_file(p) == "Hello class today we study cells"


def make_pdf(path, texts):
    doc = fitz.open()
    for t in texts:
        page = doc.new_page()
        if t:
            page.insert_text((72, 100), t, fontsize=28)
    doc.save(path)
    doc.close()


def test_pdf_text_layer(tmp_path):
    make_pdf(tmp_path / "a.pdf", ["Photosynthesis converts light into chemical energy", "Second page of notes here"])
    text = convert_file(tmp_path / "a.pdf")
    assert "[Page 1]" in text and "Photosynthesis" in text and "[Page 2]" in text


def test_image_and_scanned_pdf_ocr(tmp_path):
    try:
        from app.convert.ocr import _get_engine
        _get_engine()
    except RuntimeError as e:
        pytest.skip(str(e))
    src = tmp_path / "src.pdf"
    make_pdf(src, ["Hello Homework Helper"])
    with fitz.open(src) as d:
        d[0].get_pixmap(dpi=150).save(tmp_path / "board.png")
    assert "Homework" in convert_file(tmp_path / "board.png")
    # a PDF whose only content is an image has no text layer -> OCR fallback
    scanned = fitz.open()
    page = scanned.new_page()
    page.insert_image(page.rect, filename=str(tmp_path / "board.png"))
    scanned.save(tmp_path / "scan.pdf")
    assert "Homework" in convert_file(tmp_path / "scan.pdf")


def test_unsupported(tmp_path):
    (tmp_path / "x.xyz").write_text("hi")
    with pytest.raises(ValueError):
        convert_file(tmp_path / "x.xyz")
