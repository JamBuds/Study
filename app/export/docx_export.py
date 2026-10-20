"""Export Markdown-ish text to .docx / .md / .txt."""
import re
from pathlib import Path

import docx

_INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")


def _add_runs(paragraph, text: str) -> None:
    for part in _INLINE.split(text):
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        elif part:
            paragraph.add_run(part)


def markdown_to_docx(title: str, markdown: str, path) -> None:
    d = docx.Document()
    if title:
        d.add_heading(title, 0)
    for raw in markdown.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", line)
        b = re.match(r"^\s*[-*•]\s+(.*)", line)
        n = re.match(r"^\s*\d+[.)]\s+(.*)", line)
        if m:
            d.add_heading(m.group(2).strip("* "), min(len(m.group(1)), 4))
        elif b:
            _add_runs(d.add_paragraph(style="List Bullet"), b.group(1))
        elif n:
            _add_runs(d.add_paragraph(style="List Number"), n.group(1))
        else:
            _add_runs(d.add_paragraph(), line)
    d.save(str(path))


def export_text(title: str, text: str, path) -> None:
    path = Path(path)
    if path.suffix.lower() == ".docx":
        markdown_to_docx(title, text, path)
    else:
        path.write_text(text, encoding="utf-8")
