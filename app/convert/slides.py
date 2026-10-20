"""PowerPoint -> plain text (slide text, tables, speaker notes)."""
from pptx import Presentation
from pptx.util import Emu  # noqa: F401  (ensures the full package is importable when frozen)


def _shape_text(shape) -> list[str]:
    out = []
    if shape.has_text_frame:
        out += [p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()]
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                out.append(" | ".join(cells))
    if shape.shape_type == 6:  # group
        for sub in shape.shapes:
            out += _shape_text(sub)
    return out


def convert_pptx(path) -> str:
    prs = Presentation(str(path))
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        lines = []
        for shape in slide.shapes:
            lines += _shape_text(shape)
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        block = [f"[Slide {i}]"] + lines
        if notes:
            block += ["Speaker notes:", notes]
        parts.append("\n".join(block))
    return "\n\n".join(parts)
