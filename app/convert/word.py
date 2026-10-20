"""Word (.docx) -> plain text."""
import docx  # python-docx


def convert_docx(path) -> str:
    d = docx.Document(str(path))
    parts = [p.text.strip() for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)
